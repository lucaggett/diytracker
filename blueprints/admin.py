import io
import os
from datetime import datetime, timedelta

from flask import Blueprint, abort, current_app, flash, jsonify, redirect, render_template, request, send_file, url_for
from sqlalchemy import func
from werkzeug.utils import secure_filename

from forms import DeleteEventForm, DeleteVenueForm, EventEditForm, VenueForm, get_canton_choices
from models import db, Event, Venue
from services.auth import admin_required
from services.calendar_image import generate_weekly_calendar_image
from services.i18n import gettext as _
from services.scraper import SCRAPE_INTERVAL_HOURS, get_last_scrape_time, get_progress, is_running
from services.uploads import UPLOAD_FOLDER, allowed_file, validate_image_content

bp = Blueprint('admin', __name__)


def _fmt_duration(td):
    total = int(abs(td.total_seconds()))
    if total < 60:
        return f"{total}s"
    if total < 3600:
        return f"{total // 60}m"
    hours = total // 3600
    minutes = (total % 3600) // 60
    return f"{hours}h {minutes}m" if minutes else f"{hours}h"


@bp.route('/admin', methods=['GET'])
@admin_required
def admin():
    events = Event.query.order_by(Event.date.asc()).all()
    delete_form = DeleteEventForm()
    last_scrape = get_last_scrape_time()
    now = datetime.now()
    next_scrape = (last_scrape + timedelta(hours=SCRAPE_INTERVAL_HOURS)) if last_scrape else None
    time_since_last = _fmt_duration(now - last_scrape) + " ago" if last_scrape else "never"
    time_until_next = _fmt_duration(next_scrape - now) if next_scrape and next_scrape > now else "soon"
    today = now.date()
    current_monday = (today - timedelta(days=today.weekday())).isoformat()
    return render_template(
        'admin.html', events=events, delete_form=delete_form,
        last_scrape=last_scrape, next_scrape=next_scrape,
        time_since_last=time_since_last, time_until_next=time_until_next,
        scrape_running=is_running(), current_monday=current_monday,
    )


@bp.route('/edit_event/<int:event_id>', methods=['GET', 'POST'])
@admin_required
def edit_event(event_id):
    event = Event.query.get_or_404(event_id)
    form = EventEditForm(obj=event)
    if request.method == 'GET':
        form.venue_id.data = str(event.venue_id)
        form.venue_name.data = event.venue.name
        form.venue_address.data = event.venue.address
        form.venue_city.data = event.venue.city
        form.venue_canton.data = event.venue.canton
        form.venue_plz.data = event.venue.plz
        form.venue_coords.data = event.venue.coords or ''
        form.genre.data = [g.strip() for g in (event.genre or '').split(',') if g.strip()]
        form.end_date.data = event.end_date
        form.is_festival.data = event.is_festival
    if request.method == 'POST':
        if form.validate_on_submit():
            event.name = form.name.data
            event.date = form.date.data
            event.end_date = form.end_date.data
            event.is_festival = form.is_festival.data
            event.doors = form.doors.data
            event.acts = form.acts.data
            event.ticket_price = form.ticket_price.data
            event.ticket_link = form.ticket_link.data
            event.genre = ', '.join(form.genre.data) if form.genre.data else ''

            venue_id = form.venue_id.data
            if venue_id and venue_id != 'new':
                venue = Venue.query.get(venue_id)
                if not venue:
                    flash(_('Selected venue does not exist.'))
                    return redirect(url_for('admin.edit_event', event_id=event_id))
            else:
                venue = Venue.query.filter_by(
                    name=form.venue_name.data, city=form.venue_city.data, plz=form.venue_plz.data,
                ).first()
                if not venue:
                    venue = Venue(
                        name=form.venue_name.data,
                        address=form.venue_address.data,
                        city=form.venue_city.data,
                        canton=form.venue_canton.data,
                        plz=form.venue_plz.data,
                        coords=form.venue_coords.data,
                    )
                    db.session.add(venue)
                    db.session.flush()
            event.venue_id = venue.id

            if form.flyer.data:
                file = form.flyer.data
                if file and allowed_file(file.filename) and validate_image_content(file):
                    filename = secure_filename(file.filename)
                    flyer = os.path.join(current_app.config.get('UPLOAD_FOLDER', UPLOAD_FOLDER), filename)
                    file.save(flyer)
                    event.flyer = flyer

            db.session.commit()
            flash(_('Event updated successfully!'))
            return redirect(url_for('admin.admin'))

    return render_template('edit_event.html', form=form, event=event)


@bp.route('/delete_event/<int:event_id>', methods=['POST'])
@admin_required
def delete_event(event_id):
    form = DeleteEventForm()
    if not form.validate_on_submit():
        abort(400)
    event = Event.query.get_or_404(event_id)
    db.session.delete(event)
    db.session.commit()
    flash(_('Event deleted successfully!'))
    return redirect(url_for('admin.admin'))


@bp.route('/admin/venues', methods=['GET'])
@admin_required
def venues():
    rows = (
        db.session.query(Venue, func.count(Event.id).label('event_count'))
        .outerjoin(Event, Event.venue_id == Venue.id)
        .group_by(Venue.id)
        .order_by(func.lower(Venue.name).asc())
        .all()
    )
    return render_template('venues.html', venues=rows, delete_form=DeleteVenueForm())


@bp.route('/admin/venues/<int:venue_id>/edit', methods=['GET', 'POST'])
@admin_required
def edit_venue(venue_id):
    venue = Venue.query.get_or_404(venue_id)
    form = VenueForm(obj=venue)
    form.canton.choices = [('', _('— none —'))] + get_canton_choices()
    if request.method == 'POST' and form.validate_on_submit():
        venue.name = form.name.data.strip()
        venue.address = (form.address.data or '').strip() or None
        venue.city = form.city.data.strip()
        venue.canton = form.canton.data or None
        venue.plz = form.plz.data.strip()
        venue.coords = (form.coords.data or '').strip() or None
        db.session.commit()
        flash(_('Venue updated successfully!'))
        return redirect(url_for('admin.venues'))
    return render_template('edit_venue.html', form=form, venue=venue)


@bp.route('/admin/venues/<int:venue_id>/delete', methods=['POST'])
@admin_required
def delete_venue(venue_id):
    form = DeleteVenueForm()
    if not form.validate_on_submit():
        abort(400)
    venue = Venue.query.get_or_404(venue_id)
    event_count = Event.query.filter_by(venue_id=venue.id).count()
    if event_count > 0:
        sample = (
            Event.query.filter_by(venue_id=venue.id)
            .order_by(Event.date.asc()).limit(3).all()
        )
        names = ', '.join(e.name for e in sample)
        more = _(' and %(k)d more', k=event_count - 3) if event_count > 3 else ''
        flash(_(
            'Cannot delete venue "%(venue)s": it still has %(n)d event(s) '
            '(%(names)s%(more)s). Reassign or delete those events first.',
            venue=venue.name, n=event_count, names=names, more=more,
        ))
        return redirect(url_for('admin.venues'))
    db.session.delete(venue)
    db.session.commit()
    flash(_('Venue deleted.'))
    return redirect(url_for('admin.venues'))


@bp.route('/admin/scrape-status')
@admin_required
def scrape_status():
    if not is_running():
        return jsonify({'running': False})
    progress = get_progress()
    total = progress.get('total', 0)
    processed = progress.get('processed', 0)
    started_at = progress.get('started_at')
    eta_seconds = None
    if started_at and processed > 0 and total > 0:
        elapsed = (datetime.now() - datetime.fromisoformat(started_at)).total_seconds()
        rate = elapsed / processed
        remaining = total - processed
        eta_seconds = int(rate * remaining)
    return jsonify({
        'running': True,
        'phase': progress.get('phase', ''),
        'total': total,
        'processed': processed,
        'eta_seconds': eta_seconds,
    })


@bp.route('/admin/export-excel')
@admin_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font

    events = Event.query.order_by(Event.date.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Events'

    headers = ['Date', 'Name', 'Acts', 'Genre', 'Venue', 'City', 'Canton',
               'Doors', 'Ticket Price', 'Ticket Link', 'Source URL']
    bold = Font(bold=True)

    current_date = None
    row_num = 1

    for event in events:
        event_date = event.date.date() if event.date else None
        if event_date != current_date:
            if current_date is not None:
                row_num += 1
            cell = ws.cell(row=row_num, column=1, value=event_date.strftime('%A, %d %B %Y') if event_date else 'Unknown')
            cell.font = bold
            row_num += 1
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = bold
            row_num += 1
            current_date = event_date

        venue = event.venue
        ws.cell(row=row_num, column=1, value=event_date.strftime('%Y-%m-%d') if event_date else '')
        ws.cell(row=row_num, column=2, value=event.name)
        ws.cell(row=row_num, column=3, value=event.acts)
        ws.cell(row=row_num, column=4, value=event.genre)
        ws.cell(row=row_num, column=5, value=venue.name if venue else '')
        ws.cell(row=row_num, column=6, value=venue.city if venue else '')
        ws.cell(row=row_num, column=7, value=venue.canton if venue else '')
        ws.cell(row=row_num, column=8, value=event.doors.strftime('%H:%M') if event.doors else '')
        ws.cell(row=row_num, column=9, value=event.ticket_price)
        ws.cell(row=row_num, column=10, value=event.ticket_link)
        ws.cell(row=row_num, column=11, value=event.source_url)
        row_num += 1

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'events_{datetime.now().strftime("%Y%m%d")}.xlsx',
    )


@bp.route('/admin/weekly-image')
@admin_required
def weekly_calendar_image():
    week_param = request.args.get('week', '')
    try:
        ref = datetime.strptime(week_param, '%Y-%m-%d').date()
    except ValueError:
        ref = datetime.now().date()
    monday = ref - timedelta(days=ref.weekday())

    img = generate_weekly_calendar_image(monday)
    buf = io.BytesIO()
    img.save(buf, format='PNG', optimize=True)
    buf.seek(0)
    return send_file(
        buf,
        mimetype='image/png',
        as_attachment=True,
        download_name=f'events_week_{monday.strftime("%Y-%m-%d")}.png',
    )
