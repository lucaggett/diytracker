import io
from datetime import datetime, timedelta

from flask import (
    Blueprint,
    abort,
    current_app,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    url_for,
)
from sqlalchemy import func
from sqlalchemy.orm import joinedload

from diytracker.forms import (
    DeleteEventForm,
    DeleteVenueForm,
    EventEditForm,
    TogglePromoterForm,
    VenueForm,
    get_canton_choices,
)
from diytracker.models import db, Event, Label, ScrapeSuspect, Submitter, Venue
from diytracker.services.analytics import (
    REPORT_PATH,
    STATS_INTERVAL_MINUTES,
    TIMEFRAMES,
    generate_report,
    kick_stats_generation,
    read_stats,
    stats_age_seconds,
)
from diytracker.services.auth import admin_required
from diytracker.services import db_stats
from diytracker.services.event_views import event_popularity
from diytracker.services.cache import bust_cache
from diytracker.services.calendar_image import generate_weekly_calendar_image
from diytracker.services.events import clean_genre_string, resolve_venue_from_form
from diytracker.services.i18n import gettext as _
from diytracker.services.labels import all_label_choices
from diytracker.services.scraper import (
    SCRAPE_INTERVAL_HOURS,
    get_last_scrape_time,
    get_progress,
    is_running,
)
from diytracker.services.uploads import UPLOAD_FOLDER, save_flyer_file
from diytracker.utils import PARENT_GENRES_ORDER

bp = Blueprint("admin", __name__)


def _excel_safe(value):
    """Neutralise spreadsheet formula injection: user/scraper strings starting
    with a formula trigger character would otherwise execute when opened."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def _fmt_duration(td):
    total = int(abs(td.total_seconds()))
    if total < 60:
        return f"{total}s"
    if total < 3600:
        return f"{total // 60}m"
    hours = total // 3600
    minutes = (total % 3600) // 60
    return f"{hours}h {minutes}m" if minutes else f"{hours}h"


@bp.route("/admin", methods=["GET"])
@admin_required
def admin():
    now = datetime.now()
    today_start = datetime.combine(now.date(), datetime.min.time())
    events = (
        Event.query.filter(Event.date >= today_start).order_by(Event.date.asc()).all()
    )
    delete_form = DeleteEventForm()
    last_scrape = get_last_scrape_time()
    next_scrape = (
        (last_scrape + timedelta(hours=SCRAPE_INTERVAL_HOURS)) if last_scrape else None
    )
    time_since_last = (
        _fmt_duration(now - last_scrape) + " ago" if last_scrape else "never"
    )
    time_until_next = (
        _fmt_duration(next_scrape - now)
        if next_scrape and next_scrape > now
        else "soon"
    )
    today = now.date()
    current_monday = (today - timedelta(days=today.weekday())).isoformat()
    return render_template(
        "admin.html",
        events=events,
        delete_form=delete_form,
        last_scrape=last_scrape,
        next_scrape=next_scrape,
        time_since_last=time_since_last,
        time_until_next=time_until_next,
        scrape_running=is_running(),
        current_monday=current_monday,
        parent_genres_order=PARENT_GENRES_ORDER,
        analytics_stats=read_stats(),
        stats_interval_minutes=STATS_INTERVAL_MINUTES,
    )


@bp.route("/edit_event/<int:event_id>")
def edit_event_legacy(event_id):
    # Pre-0.31 URL; kept for old bookmarks.
    return redirect(url_for("admin.edit_event", event_id=event_id), code=301)


@bp.route("/admin/edit_event/<int:event_id>", methods=["GET", "POST"])
@admin_required
def edit_event(event_id):
    event = Event.query.get_or_404(event_id)
    form = EventEditForm(obj=event)
    form.label_id.choices = all_label_choices()
    if request.method == "GET":
        form.label_id.data = str(event.label_id) if event.label_id else ""
        form.venue_id.data = str(event.venue_id)
        form.venue_name.data = event.venue.name
        form.venue_address.data = event.venue.address
        form.venue_city.data = event.venue.city
        form.venue_canton.data = event.venue.canton
        form.venue_plz.data = event.venue.plz
        form.venue_coords.data = event.venue.coords or ""
        form.genre.data = [
            g.strip() for g in (event.genre or "").split(",") if g.strip()
        ]
        form.end_date.data = event.end_date
        form.is_festival.data = event.is_festival
    if request.method == "POST":
        if form.validate_on_submit():
            event.name = form.name.data
            event.date = form.date.data
            event.end_date = form.end_date.data
            event.is_festival = form.is_festival.data
            event.doors = form.doors.data
            event.acts = form.acts.data
            event.ticket_price = form.ticket_price.data
            event.ticket_link = form.ticket_link.data
            event.status = form.status.data
            event.genre = clean_genre_string(form.genre.data or [])
            event.label_id = int(form.label_id.data) if form.label_id.data else None

            venue, _created, error = resolve_venue_from_form(form)
            if error:
                flash(_("Selected venue does not exist."))
                return redirect(url_for("admin.edit_event", event_id=event_id))
            event.venue_id = venue.id

            saved = save_flyer_file(
                form.flyer.data, current_app.config.get("UPLOAD_FOLDER", UPLOAD_FOLDER)
            )
            if saved:
                event.flyer = saved

            db.session.commit()
            bust_cache()
            flash(_("Event updated successfully!"))
            return redirect(url_for("admin.admin"))

    return render_template("edit_event.html", form=form, event=event)


@bp.route("/admin/delete_event/<int:event_id>", methods=["POST"])
@admin_required
def delete_event(event_id):
    form = DeleteEventForm()
    if not form.validate_on_submit():
        abort(400)
    event = Event.query.get_or_404(event_id)
    db.session.delete(event)
    db.session.commit()
    bust_cache()
    flash(_("Event deleted successfully!"))
    return redirect(url_for("admin.admin"))


@bp.route("/admin/venues", methods=["GET"])
@admin_required
def venues():
    rows = (
        db.session.query(Venue, func.count(Event.id).label("event_count"))
        .outerjoin(Event, Event.venue_id == Venue.id)
        .group_by(Venue.id)
        .order_by(func.lower(Venue.name).asc())
        .all()
    )
    return render_template("venues.html", venues=rows, delete_form=DeleteVenueForm())


@bp.route("/admin/venues/<int:venue_id>/edit", methods=["GET", "POST"])
@admin_required
def edit_venue(venue_id):
    venue = Venue.query.get_or_404(venue_id)
    form = VenueForm(obj=venue)
    form.canton.choices = [("", _("— none —"))] + get_canton_choices()
    if request.method == "POST" and form.validate_on_submit():
        venue.name = form.name.data.strip()
        venue.address = (form.address.data or "").strip() or None
        venue.city = form.city.data.strip()
        venue.canton = form.canton.data or None
        venue.plz = form.plz.data.strip()
        venue.coords = (form.coords.data or "").strip() or None
        db.session.commit()
        flash(_("Venue updated successfully!"))
        return redirect(url_for("admin.venues"))
    return render_template("edit_venue.html", form=form, venue=venue)


@bp.route("/admin/venues/<int:venue_id>/delete", methods=["POST"])
@admin_required
def delete_venue(venue_id):
    form = DeleteVenueForm()
    if not form.validate_on_submit():
        abort(400)
    venue = Venue.query.get_or_404(venue_id)
    event_count = Event.query.filter_by(venue_id=venue.id).count()
    if event_count > 0:
        sample = (
            Event.query.filter_by(venue_id=venue.id)
            .order_by(Event.date.asc())
            .limit(3)
            .all()
        )
        names = ", ".join(e.name for e in sample)
        more = _(" and %(k)d more", k=event_count - 3) if event_count > 3 else ""
        flash(
            _(
                'Cannot delete venue "%(venue)s": it still has %(n)d event(s) '
                "(%(names)s%(more)s). Reassign or delete those events first.",
                venue=venue.name,
                n=event_count,
                names=names,
                more=more,
            )
        )
        return redirect(url_for("admin.venues"))
    db.session.delete(venue)
    db.session.commit()
    flash(_("Venue deleted."))
    return redirect(url_for("admin.venues"))


@bp.route("/admin/venues/<int:venue_id>/accessibility-link")
@admin_required
def generate_accessibility_link(venue_id):
    venue = Venue.query.get_or_404(venue_id)
    venue.generate_accessibility_token()
    db.session.commit()
    link = url_for(
        "public.accessibility_form", token=venue.accessibility_token, _external=True
    )
    flash(link, "accessibility_link")
    return redirect(url_for("admin.venues"))


@bp.route("/admin/users", methods=["GET"])
@admin_required
def users():
    all_users = Submitter.query.order_by(Submitter.email.asc()).all()
    label_counts = dict(
        db.session.query(Label.promoter_id, func.count(Label.id)).group_by(
            Label.promoter_id
        )
    )
    return render_template(
        "admin_users.html",
        users=all_users,
        label_counts=label_counts,
        toggle_form=TogglePromoterForm(),
    )


@bp.route("/admin/users/<int:user_id>/toggle-promoter", methods=["POST"])
@admin_required
def toggle_promoter(user_id):
    form = TogglePromoterForm()
    if not form.validate_on_submit():
        abort(400)
    user = Submitter.query.get_or_404(user_id)
    user.is_promoter = not user.is_promoter
    db.session.commit()
    status = _("granted") if user.is_promoter else _("revoked")
    flash(
        _("Promoter status %(status)s for %(email)s.", status=status, email=user.email)
    )
    return redirect(url_for("admin.users"))


@bp.route("/admin/scrape-status")
@admin_required
def scrape_status():
    if not is_running():
        return jsonify({"running": False})
    progress = get_progress()
    total = progress.get("total", 0)
    processed = progress.get("processed", 0)
    started_at = progress.get("started_at")
    eta_seconds = None
    if started_at and processed > 0 and total > 0:
        elapsed = (datetime.now() - datetime.fromisoformat(started_at)).total_seconds()
        rate = elapsed / processed
        remaining = total - processed
        eta_seconds = int(rate * remaining)
    return jsonify(
        {
            "running": True,
            "phase": progress.get("phase", ""),
            "total": total,
            "processed": processed,
            "eta_seconds": eta_seconds,
        }
    )


@bp.route("/admin/export-excel")
@admin_required
def export_excel():
    import openpyxl
    from openpyxl.styles import Font

    events = (
        Event.query.options(joinedload(Event.venue)).order_by(Event.date.asc()).all()
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Events"

    headers = [
        "Date",
        "Name",
        "Acts",
        "Genre",
        "Venue",
        "City",
        "Canton",
        "Doors",
        "Ticket Price",
        "Ticket Link",
        "Source URL",
    ]
    bold = Font(bold=True)

    current_date = None
    row_num = 1

    for event in events:
        event_date = event.date.date() if event.date else None
        if event_date != current_date:
            if current_date is not None:
                row_num += 1
            cell = ws.cell(
                row=row_num,
                column=1,
                value=event_date.strftime("%A, %d %B %Y") if event_date else "Unknown",
            )
            cell.font = bold
            row_num += 1
            for col, h in enumerate(headers, 1):
                cell = ws.cell(row=row_num, column=col, value=h)
                cell.font = bold
            row_num += 1
            current_date = event_date

        venue = event.venue
        ws.cell(
            row=row_num,
            column=1,
            value=event_date.strftime("%Y-%m-%d") if event_date else "",
        )
        ws.cell(row=row_num, column=2, value=_excel_safe(event.name))
        ws.cell(row=row_num, column=3, value=_excel_safe(event.acts))
        ws.cell(row=row_num, column=4, value=_excel_safe(event.genre))
        ws.cell(row=row_num, column=5, value=_excel_safe(venue.name) if venue else "")
        ws.cell(row=row_num, column=6, value=_excel_safe(venue.city) if venue else "")
        ws.cell(row=row_num, column=7, value=_excel_safe(venue.canton) if venue else "")
        ws.cell(
            row=row_num,
            column=8,
            value=event.doors.strftime("%H:%M") if event.doors else "",
        )
        ws.cell(row=row_num, column=9, value=_excel_safe(event.ticket_price))
        ws.cell(row=row_num, column=10, value=_excel_safe(event.ticket_link))
        ws.cell(row=row_num, column=11, value=_excel_safe(event.source_url))
        row_num += 1

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"events_{datetime.now().strftime('%Y%m%d')}.xlsx",
    )


@bp.route("/admin/analytics", methods=["GET", "POST"])
@admin_required
def analytics():
    if request.method == "POST":
        timeframe = request.form.get("timeframe", "7d")
        valid = {t for t, _ in TIMEFRAMES}
        if timeframe not in valid:
            timeframe = "7d"
        success, message = generate_report(timeframe)
        flash(message, "success" if success else "error")
        return redirect(url_for("admin.analytics"))
    report_exists = REPORT_PATH.exists()
    report_mtime = (
        datetime.fromtimestamp(REPORT_PATH.stat().st_mtime) if report_exists else None
    )
    return render_template(
        "analytics.html",
        timeframes=TIMEFRAMES,
        report_exists=report_exists,
        report_mtime=report_mtime,
    )


@bp.route("/admin/analytics/view")
@admin_required
def analytics_view():
    if not REPORT_PATH.exists():
        flash(_("No report has been generated yet. Use the form to create one."))
        return redirect(url_for("admin.analytics"))
    return send_file(REPORT_PATH, mimetype="text/html")


@bp.route("/admin/analytics/stats")
@admin_required
def analytics_stats():
    stats = read_stats()
    age = stats_age_seconds()
    # Regenerate in the background when the cache is missing or the scheduler
    # looks dead (dev mode, or the first gunicorn worker died). Never generate
    # inline: goaccess over 60 days of logs must not block a request.
    if stats is None or age is None or age > 3 * STATS_INTERVAL_MINUTES * 60:
        kick_stats_generation(current_app._get_current_object())
    if stats is None:
        return jsonify({"status": "pending"}), 202
    return jsonify({"status": "ok", "age_seconds": age, **stats})


@bp.route("/admin/statistics")
@admin_required
def statistics():
    now = datetime.now()
    return render_template(
        "admin_statistics.html",
        popularity=event_popularity(),
        monthly=db_stats.monthly_series(),
        yearly=db_stats.events_per_year(),
        by_genre=db_stats.events_by_parent_genre(),
        by_canton=db_stats.events_by_canton(),
        top_venues=db_stats.top_venues(),
        price_stats=db_stats.ticket_price_stats(),
        now=now,
    )


@bp.route("/admin/scrape-suspects", methods=["GET"])
@admin_required
def scrape_suspects():
    import json

    suspects = ScrapeSuspect.query.order_by(ScrapeSuspect.last_seen.desc()).all()
    rows = [
        {
            "suspect": s,
            "signals": json.loads(s.signals or "[]"),
            "sample_paths": json.loads(s.sample_paths or "[]"),
        }
        for s in suspects
    ]
    return render_template("scrape_suspects.html", rows=rows)


@bp.route("/admin/scrape-suspects/clear", methods=["POST"])
@admin_required
def clear_scrape_suspects():
    deleted = ScrapeSuspect.query.delete()
    db.session.commit()
    flash(_("Cleared %(n)d suspect(s).", n=deleted), "success")
    return redirect(url_for("admin.scrape_suspects"))


@bp.route("/admin/weekly-image")
@admin_required
def weekly_calendar_image():
    week_param = request.args.get("week", "")
    try:
        ref = datetime.strptime(week_param, "%Y-%m-%d").date()
    except ValueError:
        ref = datetime.now().date()
    monday = ref - timedelta(days=ref.weekday())

    genre_param = (request.args.get("genre") or "").strip() or None
    if genre_param not in PARENT_GENRES_ORDER:
        genre_param = None

    img = generate_weekly_calendar_image(monday, parent_genre=genre_param)
    buf = io.BytesIO()
    img.save(buf, format="PNG", optimize=True)
    buf.seek(0)
    suffix = f"_{genre_param.lower().replace('/', '-')}" if genre_param else ""
    return send_file(
        buf,
        mimetype="image/png",
        as_attachment=True,
        download_name=f"events_week_{monday.strftime('%Y-%m-%d')}{suffix}.png",
    )
